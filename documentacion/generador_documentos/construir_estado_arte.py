# -*- coding: utf-8 -*-
"""
Construye la Tarea 3 de Redacción de Investigación de Textos Académicos (estado del arte) sobre una copia
del documento base de la UNEMI: conserva la portada institucional, el encabezado, el pie, la tabla de
contenido y los índices; quita los certificados, la declaración, el resumen y el anexo académico, que no
pide esta tarea; y escribe el contenido de _tarea3_contenido.py con los mismos estilos del trabajo final.

Uso:
    python documentacion/generador_documentos/construir_estado_arte.py \\
        --base AVANCE1_MARLON_VICTOR.docx --salida documentacion/avances/LluilemaPisco_Victor_Tarea3_EstadoDelArte.docx

Autores: Marlon David Pérez Almachi, Víctor Manuel Lluilema Pisco
Proyecto de titulación · Universidad Estatal de Milagro · 2026
"""
from __future__ import annotations

import argparse
import copy
import re
import shutil
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import _tarea3_contenido as contenido
from construir_trabajo_final import Constructor, _buscar, _poner_texto, indices_a_figura

_REF = re.compile(r'\{([FT]):([a-z0-9_]+)\}')
ASIGNATURA = 'Redacción de Investigación de Textos Académicos'
DOCENTE = 'Msc. Espinoza Bravo María Gabriela'
FECHA = 'Milagro, octubre de 2026'


def numerar(crudos: list[tuple]) -> list[tuple]:
    """Asigna números a figuras y tablas en orden de aparición y resuelve {F:clave} y {T:clave}."""
    numeros = {'F': {}, 'T': {}}
    for bq in crudos:
        if bq[0] in ('fig', 'tab'):
            t = 'F' if bq[0] == 'fig' else 'T'
            numeros[t][bq[1]] = len(numeros[t]) + 1

    def resolver(texto: str) -> str:
        return _REF.sub(lambda m: f"{'Figura' if m.group(1) == 'F' else 'Tabla'} {numeros[m.group(1)][m.group(2)]}",
                        texto)
    salida = []
    for bq in crudos:
        if bq[0] == 'p':
            salida.append(('p', resolver(bq[1])))
        elif bq[0] == 'ul':
            salida.append(('ul', [resolver(x) for x in bq[1]]))
        elif bq[0] in ('fig', 'tab'):
            salida.append((bq[0],) + tuple(bq[2:]))
        else:
            salida.append(bq)
    return salida


def _tabla_portada(doc):
    return next(t for t in doc.tables if any(r.cells[0].text.strip().startswith('Facultad') for r in t.rows))


def portada(doc) -> None:
    """Portada institucional de la tarea: tema, asignatura, docente, estudiantes, carpeta y fecha."""
    for p in doc.paragraphs[:10]:
        t = p.text.strip()
        if t.startswith('ESTUDIO DE CASO PREVIO'):
            _poner_texto(p, ASIGNATURA.upper())
        elif t.startswith('DE: INGENIERO'):
            _poner_texto(p, 'TAREA 3: ESTADO DEL ARTE DEL TRABAJO DE TITULACIÓN')
    tabla = _tabla_portada(doc)
    molde = tabla.rows[0]._tr
    fila_fecha = next((r._tr for r in tabla.rows if r.cells[0].text.strip().startswith('Milagro')
                       or r.cells[-1].text.strip().startswith('Milagro')), None)

    def fila(etiqueta, valor, antes):
        tr = copy.deepcopy(molde)
        for tc, texto in zip(tr.findall(qn('w:tc')), (etiqueta, valor)):
            ps = tc.findall(qn('w:p'))
            for extra in ps[1:]:
                tc.remove(extra)
            _poner_texto(Paragraph(ps[0], tabla._parent), texto)
        if antes is not None:
            antes.addprevious(tr)
        else:
            tabla._tbl.append(tr)

    fila('Tema:', contenido.TITULO, molde)
    fila('Asignatura:', ASIGNATURA, fila_fecha)
    fila('Carpeta digital:', contenido.DRIVE, fila_fecha)
    for row in tabla.rows:
        celdas = row.cells
        if celdas[0].text.strip().startswith('Tutor'):
            _poner_texto(celdas[0].paragraphs[0], 'Docente:')
            for extra in celdas[1].paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            _poner_texto(celdas[1].paragraphs[0], DOCENTE)
    for p_el in doc.element.body.iter(qn('w:p')):
        texto = ''.join(t.text or '' for t in p_el.iter(qn('w:t')))
        if texto.strip().startswith('Milagro,') and '2026' in texto:
            _poner_texto(Paragraph(p_el, doc._body), FECHA)


def _vacio(el) -> bool:
    return el.tag == qn('w:p') and not ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip() \
        and el.find('.//' + qn('w:drawing')) is None and el.find('.//' + qn('w:sectPr')) is None \
        and el.find('.//' + qn('w:fldChar')) is None      # nunca romper un campo (fin del índice)


def quitar_preliminares(doc) -> None:
    """Quita la ficha, los certificados, la declaración, el enlace del Avance 1, el resumen y el abstract."""
    body = doc.element.body
    tabla_portada = _tabla_portada(doc)._tbl
    toc = _buscar(doc, 'TABLA DE CONTENIDO')._p
    el = tabla_portada.getnext()
    while el is not None and el is not toc:
        sig = el.getnext()
        body.remove(el)
        el = sig
    ppr = toc.get_or_add_pPr()
    if ppr.find(qn('w:pageBreakBefore')) is None:
        ppr.insert(1 if ppr.find(qn('w:pStyle')) is not None else 0, OxmlElement('w:pageBreakBefore'))
    for marca in ('ENLACE DEL VIDEO', 'RESUMEN', 'Se desarrolló un prototipo', 'Palabras clave', 'ABSTRACT',
                  'A business intelligence prototype', 'Keywords'):
        try:
            p = _buscar(doc, marca)
            p._p.getparent().remove(p._p)
        except LookupError:
            pass


def vaciar(doc):
    """Quita el cuerpo y el anexo académico; devuelve el elemento tras el cual se escribe."""
    inicio = _buscar(doc, '1. Introducción', 'Heading 1')._p
    previo = inicio.getprevious()
    while previo is not None and _vacio(previo):     # páginas en blanco que dejaron el resumen y el abstract
        anterior = previo.getprevious()
        previo.getparent().remove(previo)
        previo = anterior
    el = inicio
    while el is not None:
        sig = el.getnext()
        if el.tag == qn('w:sectPr'):
            break
        el.getparent().remove(el)
        el = sig
    return previo


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--base', required=True)
    ap.add_argument('--salida', required=True)
    a = ap.parse_args()
    salida = Path(a.salida)
    shutil.copyfile(a.base, salida)
    doc = docx.Document(str(salida))
    c = Constructor(doc)
    c.tomar_modelos()
    c.salto_capitulo = False      # tarea breve: las secciones continúan en la misma página
    portada(doc)
    quitar_preliminares(doc)
    c.ancla = vaciar(doc)
    primero = True
    for bq in numerar(contenido.bloques()):
        tipo = bq[0]
        if tipo in ('h1', 'h2', 'h3'):
            el = c.titulo(int(tipo[1]), bq[1])
            if primero or bq[1].startswith('13. Referencias'):   # el cuerpo y las referencias en página nueva
                ppr = el.find(qn('w:pPr'))
                if ppr.find(qn('w:pageBreakBefore')) is None:
                    ppr.insert(1, OxmlElement('w:pageBreakBefore'))
                primero = False
        elif tipo == 'p':
            c.parrafo(bq[1])
        elif tipo == 'ul':
            c.vinetas(bq[1])
        elif tipo == 'fig':
            c.figura(*bq[1:])
        elif tipo == 'tab':
            c.tabla(*bq[1:])
        elif tipo == 'ref':
            c.referencia(bq[1])
    indices_a_figura(doc)
    doc.save(str(salida))
    print(f'{salida}: {c.n_fig} figuras, {c.n_tab} tablas')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
