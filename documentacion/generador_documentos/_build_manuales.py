# -*- coding: utf-8 -*-
"""
Generador de los manuales de entrega en Word.

Convierte los manuales de documentacion/manuales/*.md (la fuente única) en
documentos .docx con portada institucional, índice, encabezado y numeración de
páginas. El PDF se obtiene después con _docx_a_pdf.ps1, que actualiza el índice.
Solo interpreta el subconjunto de Markdown que usan los manuales: títulos,
párrafos, negritas, cursivas, código, listas, tablas, imágenes, citas y
bloques de código.

Uso:
    python _build_manuales.py            # genera los tres manuales
    python _build_manuales.py MANUAL_DE_USUARIO

Autores: Marlon David Pérez Almachi, Víctor Manuel Lluilema Pisco
Proyecto de titulación · Universidad Estatal de Milagro · 2026
"""
from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[2]
CARPETA = RAIZ / "documentacion" / "manuales"
LOGO = RAIZ / "codigo_fuente" / "frontend" / "public" / "logo.png"
AUTORES = ["Marlon David Pérez Almachi", "Víctor Manuel Lluilema Pisco"]
MANUALES = ["MANUAL_DE_USUARIO", "MANUAL_TECNICO_DE_INSTALACION", "GUIA_DE_PRUEBAS_DE_ACEPTACION"]

AZUL = RGBColor(0x00, 0x2E, 0x45)
NARANJA = RGBColor(0xFF, 0x69, 0x00)
GRIS = RGBColor(0x5B, 0x73, 0x85)
AVISOS = {"NOTE": ("Nota", "E8F0F5"), "TIP": ("Consejo", "E6F4EA"), "IMPORTANT": ("Importante", "FFF1E6"),
          "WARNING": ("Advertencia", "FDECEA")}
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")


def sombrear(elemento, color: str) -> None:
    """Aplica un color de fondo a un párrafo o a una celda."""
    propiedades = elemento.get_or_add_pPr() if hasattr(elemento, "get_or_add_pPr") else elemento.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    propiedades.append(sombra)


def texto_enriquecido(parrafo, texto: str, tam: float | None = None, color: RGBColor | None = None) -> None:
    """Agrega el texto con negritas, cursivas, código y enlaces de Markdown como formatos de Word."""
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        negrita = cursiva = codigo = False
        if trozo.startswith("**") and trozo.endswith("**"):
            trozo, negrita = trozo[2:-2], True
        elif trozo.startswith("`") and trozo.endswith("`"):
            trozo, codigo = trozo[1:-1], True
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            trozo, cursiva = trozo[1:-1], True
        elif trozo.startswith("[") and "](" in trozo:
            etiqueta, destino = trozo[1:-1].split("](", 1)
            trozo = etiqueta + (f" ({destino})" if destino.startswith("http") else "")
        elif trozo.startswith("<http"):
            trozo = trozo[1:-1]
        run = parrafo.add_run(trozo)
        run.bold, run.italic = negrita or None, cursiva or None
        if codigo:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = AZUL
        elif tam:
            run.font.size = Pt(tam)
        if color and not codigo:
            run.font.color.rgb = color


def campo(parrafo, instruccion: str) -> None:
    """Inserta un campo de Word (número de página, índice) que Word calcula al abrir o exportar."""
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), ("end", None)):
        run = parrafo.add_run()
        if tipo:
            marca = OxmlElement("w:fldChar")
            marca.set(qn("w:fldCharType"), tipo)
            run._r.append(marca)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run._r.append(instr)


def preparar_estilos(doc: Document) -> None:
    """Define la tipografía y los colores de los estilos que usan los manuales."""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for nivel, tam in ((1, 18), (2, 14), (3, 12)):
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tam)
        estilo.font.bold = True
        estilo.font.color.rgb = AZUL if nivel < 3 else NARANJA
        estilo.paragraph_format.space_before = Pt(16 if nivel == 1 else 12)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True


def portada(doc: Document, titulo: str, meta: dict) -> None:
    """Primera página: logo, universidad, título del manual, sistema, autores y versión."""
    centro = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = centro
    p.paragraph_format.space_before = Pt(36)
    p.add_run().add_picture(str(LOGO), width=Cm(4.2))
    for texto, tam, color, negrita in (
        ("UNIVERSIDAD ESTATAL DE MILAGRO", 15, AZUL, True),
        ("Facultad de Ciencias e Ingenierías", 12, GRIS, False),
        ("Carrera de Ingeniería en Tecnologías de la Información", 12, GRIS, False),
    ):
        p = doc.add_paragraph(); p.alignment = centro
        r = p.add_run(texto); r.font.size = Pt(tam); r.font.color.rgb = color; r.bold = negrita
    p = doc.add_paragraph(); p.alignment = centro; p.paragraph_format.space_before = Pt(60)
    r = p.add_run(titulo.upper()); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = AZUL
    p = doc.add_paragraph(); p.alignment = centro
    r = p.add_run("VicMar Data AI"); r.font.size = Pt(18); r.bold = True; r.font.color.rgb = NARANJA
    p = doc.add_paragraph(); p.alignment = centro
    r = p.add_run("Predicción del abandono estudiantil con inteligencia de negocios\ny consultas en lenguaje natural")
    r.font.size = Pt(12); r.font.color.rgb = GRIS
    p = doc.add_paragraph(); p.alignment = centro; p.paragraph_format.space_before = Pt(70)
    r = p.add_run("Autores"); r.bold = True; r.font.color.rgb = AZUL
    for autor in AUTORES:
        p = doc.add_paragraph(autor); p.alignment = centro; p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph(); p.alignment = centro; p.paragraph_format.space_before = Pt(40)
    r = p.add_run(f"{meta.get('Versión del documento', 'Versión 1.0')} · Milagro, Ecuador")
    r.font.size = Pt(10); r.font.color.rgb = GRIS
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(); r = p.add_run("Contenido"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = AZUL
    campo(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def encabezado_y_pie(doc: Document, titulo: str) -> None:
    """Encabezado con el nombre del manual y pie con la numeración de páginas."""
    seccion = doc.sections[0]
    seccion.different_first_page_header_footer = True
    p = seccion.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"VicMar Data AI · {titulo}"); r.font.size = Pt(8.5); r.font.color.rgb = GRIS
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Universidad Estatal de Milagro · Página "); r.font.size = Pt(8.5); r.font.color.rgb = GRIS
    campo(p, "PAGE")


def tabla(doc: Document, filas: list[list[str]]) -> None:
    """Tabla con encabezado azul y filas alternadas."""
    columnas = max(len(f) for f in filas)
    t = doc.add_table(rows=len(filas), cols=columnas)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j in range(columnas):
            celda = t.cell(i, j)
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            texto_enriquecido(p, fila[j] if j < len(fila) else "", tam=9.5,
                              color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                sombrear(celda._tc, "002E45")
            elif i % 2 == 0:
                sombrear(celda._tc, "F2F6F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def imagen(doc: Document, ruta_md: str, base: Path, ancho_cm: float = 15.5) -> None:
    """Inserta una imagen centrada; las capturas verticales (móvil) van más angostas."""
    ruta = (base / ruta_md).resolve()
    if not ruta.exists():
        return
    from PIL import Image
    with Image.open(ruta) as im:
        if im.height > im.width:
            ancho_cm = 6.5
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))


def aviso(doc: Document, tipo: str, lineas: list[str]) -> None:
    """Cita con etiqueta (Nota, Consejo, Importante) en un recuadro de color."""
    etiqueta, color = AVISOS.get(tipo, ("", "EEF2F5"))
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    sombrear(p._p, color)
    if etiqueta:
        r = p.add_run(etiqueta + ". "); r.bold = True; r.font.color.rgb = AZUL
    texto_enriquecido(p, " ".join(lineas))


def convertir(nombre: str) -> Path:
    """Genera el .docx de un manual a partir de su Markdown."""
    fuente = CARPETA / f"{nombre}.md"
    lineas = fuente.read_text(encoding="utf-8").splitlines()
    titulo = lineas[0].lstrip("# ").strip()
    meta = {}
    i = 1
    while i < len(lineas) and not lineas[i].startswith("---"):
        m = re.match(r"\*\*(.+?):\*\*\s*(.*)", lineas[i])
        if m:
            meta[m.group(1)] = m.group(2)
        i += 1
    para_quien = next((f"**Para quién es:**{v}" for k, v in meta.items() if k == "Para quién es"), "")

    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.2)
        s.left_margin, s.right_margin = Cm(2.5), Cm(2.5)
    preparar_estilos(doc)
    portada(doc, titulo, meta)
    encabezado_y_pie(doc, titulo)
    doc.core_properties.author = "; ".join(AUTORES)
    doc.core_properties.last_modified_by = AUTORES[0]
    doc.core_properties.title = f"{titulo} — VicMar Data AI"
    doc.core_properties.subject = "Proyecto de titulación · Universidad Estatal de Milagro"
    if para_quien:
        texto_enriquecido(doc.add_paragraph(), para_quien)

    while i < len(lineas):
        linea = lineas[i]
        sin = linea.strip()
        if not sin or sin == "---":
            i += 1; continue
        if sin.startswith("```"):
            bloque = []
            i += 1
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i]); i += 1
            p = doc.add_paragraph(); sombrear(p._p, "F2F4F6")
            p.paragraph_format.left_indent = Cm(0.3)
            r = p.add_run("\n".join(bloque)); r.font.name = "Consolas"; r.font.size = Pt(8.5)
            i += 1; continue
        m = re.match(r"(#{1,3})\s+(.*)", sin)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            i += 1; continue
        if sin.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in celdas if c):
                    filas.append(celdas)
                i += 1
            tabla(doc, filas); continue
        if sin.startswith(">"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                bloque.append(lineas[i].strip()[1:].strip()); i += 1
            tipo = ""
            if bloque and (t := re.fullmatch(r"\[!(\w+)\]", bloque[0])):
                tipo, bloque = t.group(1), bloque[1:]
            aviso(doc, tipo, bloque); continue
        img = re.search(r"!\[[^\]]*\]\(([^)]+)\)", sin) or re.search(r'<img[^>]*src="([^"]+)"', sin)
        if img:
            imagen(doc, img.group(1), CARPETA)
            i += 1; continue
        if sin.startswith("<"):
            i += 1; continue
        m = re.match(r"(\s*)([-*]|\d+\.)\s+(.*)", linea)
        if m:
            estilo = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            p = doc.add_paragraph(style=estilo)
            if m.group(1):
                p.paragraph_format.left_indent = Cm(1.6)
            texto_enriquecido(p, m.group(3))
            i += 1; continue
        partes = [sin]
        i += 1
        while i < len(lineas) and lineas[i].strip() and not re.match(r"\s*(#|\||>|```|[-*]\s|\d+\.\s|!\[|<)", lineas[i]):
            partes.append(lineas[i].strip()); i += 1
        texto_enriquecido(doc.add_paragraph(), " ".join(partes))

    destino = CARPETA / f"{nombre}.docx"
    doc.save(destino)
    return destino


def main() -> int:
    """Genera los manuales indicados (o todos) e informa dónde quedaron."""
    nombres = sys.argv[1:] or MANUALES
    for nombre in nombres:
        print("Generado:", convertir(nombre).relative_to(RAIZ))
    print(f"Listo ({date.today():%d/%m/%Y}). Para el PDF: powershell -File _docx_a_pdf.ps1 -Docx <archivo.docx>")
    return 0


if __name__ == "__main__":
    sys.exit(main())
