# -*- coding: utf-8 -*-
"""
Utilidades para insertar contenido en un .docx existente.

Insertan títulos, párrafos, viñetas, figuras y tablas en formato APA 7 y
fragmentos de código después de un ancla, respetando el formato del documento
oficial de titulación (Arial 12, justificado, interlineado 1,5, 10 pt después
de cada párrafo; títulos con Heading 1 / Heading 2). Reciben el documento como
parámetro, sin variables globales; las usa _build_avance2.py. Markdown mínimo
soportado en el texto: **negrita** y *cursiva*.

Autores: Marlon David Pérez Almachi, Víctor Manuel Lluilema Pisco
Proyecto de titulación · Universidad Estatal de Milagro · 2026
"""
from __future__ import annotations

import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

FUENTE = 'Arial'
TAM = Pt(12)
INTERLINEADO = 1.5
DESPUES = Emu(127000)  # 10 pt, igual que el cuerpo del documento oficial


class Insertor:
    """Encadena inserciones después de un ancla (párrafo o tabla).

    Cada método público inserta tras `anchor` y devuelve el último elemento
    insertado, que sirve de ancla para la inserción siguiente. Lleva la
    cuenta de figuras y tablas para numerarlas en orden.
    """

    def __init__(self, doc: docx.Document):
        self.doc = doc
        self.body = doc.paragraphs[0]._parent
        self.n_figura = 0
        self.n_tabla = 0

    # ── utilidades internas ─────────────────────────────────────────────
    @staticmethod
    def _el(anchor):
        return anchor._p if hasattr(anchor, '_p') else anchor._tbl

    def _nuevo_parrafo(self, estilo: str | None = None):
        p = docx.text.paragraph.Paragraph(OxmlElement('w:p'), self.body)
        if estilo:
            try:
                p.style = estilo
            except KeyError:
                pass
        return p

    def _formato_cuerpo(self, p, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY, despues=DESPUES):
        pf = p.paragraph_format
        pf.line_spacing = INTERLINEADO
        pf.space_after = despues
        if alineacion is not None:
            p.alignment = alineacion

    @staticmethod
    def _run(p, texto, negrita=False, cursiva=False, tam=TAM, fuente=FUENTE, color=None):
        r = p.add_run(texto)
        r.bold, r.italic = negrita, cursiva
        r.font.name = fuente
        r.font.size = tam
        r._element.rPr.rFonts.set(qn('w:eastAsia'), fuente)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return r

    def _markdown(self, p, texto, tam=TAM, negrita_base=False):
        """Agrega `texto` a `p` como runs, interpretando **negrita** y *cursiva* (sin anidar)."""
        for trozo in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', texto):
            if not trozo:
                continue
            if trozo.startswith('**') and trozo.endswith('**'):
                self._run(p, trozo[2:-2], negrita=True, tam=tam)
            elif trozo.startswith('*') and trozo.endswith('*') and len(trozo) > 2:
                self._run(p, trozo[1:-1], cursiva=True, negrita=negrita_base, tam=tam)
            else:
                self._run(p, trozo, negrita=negrita_base, tam=tam)

    def _insertar(self, anchor, nuevo):
        self._el(anchor).addnext(self._el(nuevo))
        return nuevo

    # ── API pública ─────────────────────────────────────────────────────
    def titulo(self, anchor, texto, nivel=1, salto_pagina=False):
        """Inserta un título con el estilo Heading `nivel` (14 pt en nivel 1, 12 pt en los demás).

        Con `salto_pagina`, inserta antes un salto de página.
        """
        if salto_pagina:
            anchor = self.salto_pagina(anchor)
        p = self._nuevo_parrafo(f'Heading {nivel}')
        self._run(p, texto, negrita=True, tam=Pt(14 if nivel == 1 else 12))
        p.paragraph_format.space_after = DESPUES
        if nivel > 1:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = INTERLINEADO
        return self._insertar(anchor, p)

    def subtitulo(self, anchor, texto):
        """Inserta un subtítulo en negrita dentro de una sección (mismo estilo que
        'Objetivo general' / 'Comparación crítica...' del documento oficial)."""
        p = self._nuevo_parrafo('Normal')
        self._formato_cuerpo(p)
        self._run(p, texto, negrita=True)
        p.paragraph_format.keep_with_next = True
        return self._insertar(anchor, p)

    def parrafo(self, anchor, texto, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """Inserta un párrafo de cuerpo (justificado por defecto) con Markdown mínimo."""
        p = self._nuevo_parrafo('Normal')
        self._formato_cuerpo(p, alineacion)
        self._markdown(p, texto)
        return self._insertar(anchor, p)

    def bloque(self, anchor, texto):
        """Inserta un párrafo por cada bloque separado por línea en blanco.

        Los saltos de línea simples dentro de un bloque se unen con un espacio.
        """
        for t in [b.strip() for b in texto.strip().split('\n\n') if b.strip()]:
            anchor = self.parrafo(anchor, ' '.join(t.split('\n')))
        return anchor

    def viñetas(self, anchor, items):
        """Inserta una lista con viñeta (•) y sangría francesa, un párrafo por elemento."""
        for it in items:
            p = self._nuevo_parrafo('Normal')
            self._formato_cuerpo(p, despues=Emu(50800))
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            self._run(p, '• ')
            self._markdown(p, it)
            anchor = self._insertar(anchor, p)
        return anchor

    def salto_pagina(self, anchor):
        """Inserta un párrafo que contiene solo un salto de página."""
        p = self._nuevo_parrafo('Normal')
        p.add_run().add_break(WD_BREAK.PAGE)
        return self._insertar(anchor, p)

    def leyenda(self, anchor, texto):
        """Inserta una leyenda centrada en cursiva de 10 pt."""
        p = self._nuevo_parrafo('Normal')
        self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.CENTER)
        self._run(p, texto, cursiva=True, tam=Pt(10))
        return self._insertar(anchor, p)

    def figura(self, anchor, ruta_imagen, titulo, fuente='Elaboración propia.', ancho_pulgadas=6.2):
        """Inserta una figura en formato APA 7: 'Figura N' en negrita + título en
        cursiva ARRIBA de la imagen, y la nota de fuente debajo."""
        self.n_figura += 1
        p = self._nuevo_parrafo('Normal')
        self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.LEFT, despues=Emu(0))
        self._run(p, f'Figura {self.n_figura}', negrita=True, tam=Pt(11))
        p.paragraph_format.keep_with_next = True
        anchor = self._insertar(anchor, p)
        p = self._nuevo_parrafo('Normal')
        self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.LEFT, despues=Emu(50800))
        self._run(p, titulo, cursiva=True, tam=Pt(11))
        p.paragraph_format.keep_with_next = True
        anchor = self._insertar(anchor, p)
        p = self._nuevo_parrafo('Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ruta_imagen), width=Inches(ancho_pulgadas))
        anchor = self._insertar(anchor, p)
        p = self._nuevo_parrafo('Normal')
        self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.LEFT)
        self._run(p, 'Nota. ', cursiva=True, tam=Pt(10))
        self._run(p, fuente, tam=Pt(10))
        return self._insertar(anchor, p)

    def tabla(self, anchor, filas, titulo=None, nota='Elaboración propia.', anchos=None, tam=9):
        """Inserta una tabla APA 7 ('Tabla N' + título en cursiva arriba, nota abajo).

        La primera fila es el encabezado (sombreado y repetido si la tabla
        cruza de página). Sin `titulo`, la tabla no se numera.
        """
        if titulo:
            self.n_tabla += 1
            p = self._nuevo_parrafo('Normal')
            self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.LEFT, despues=Emu(0))
            self._run(p, f'Tabla {self.n_tabla}', negrita=True, tam=Pt(11))
            p.paragraph_format.keep_with_next = True
            anchor = self._insertar(anchor, p)
            p = self._nuevo_parrafo('Normal')
            self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.LEFT, despues=Emu(50800))
            self._run(p, titulo, cursiva=True, tam=Pt(11))
            p.paragraph_format.keep_with_next = True
            anchor = self._insertar(anchor, p)
        t = self.doc.add_table(rows=len(filas), cols=len(filas[0]))
        try:
            t.style = 'Table Grid'
        except KeyError:
            pass
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, fila in enumerate(filas):
            for j, valor in enumerate(fila):
                celda = t.cell(i, j)
                celda.text = ''
                par = celda.paragraphs[0]
                par.paragraph_format.space_after = Emu(0)
                self._markdown(par, str(valor), tam=Pt(tam), negrita_base=(i == 0))
                if i == 0:
                    sombra = OxmlElement('w:shd')
                    sombra.set(qn('w:val'), 'clear')
                    sombra.set(qn('w:color'), 'auto')
                    sombra.set(qn('w:fill'), 'D9E2F3')
                    celda._tc.get_or_add_tcPr().append(sombra)
        if anchos:
            for fila in t.rows:
                for j, ancho in enumerate(anchos):
                    fila.cells[j].width = Inches(ancho)
        # Repetir encabezado si la tabla cruza de página.
        trPr = t.rows[0]._tr.get_or_add_trPr()
        rep = OxmlElement('w:tblHeader')
        rep.set(qn('w:val'), 'true')
        trPr.append(rep)
        anchor = self._insertar(anchor, t)
        if nota:
            p = self._nuevo_parrafo('Normal')
            self._formato_cuerpo(p, WD_ALIGN_PARAGRAPH.LEFT)
            self._run(p, 'Nota. ', cursiva=True, tam=Pt(10))
            self._markdown(p, nota, tam=Pt(10))
            anchor = self._insertar(anchor, p)
        return anchor

    def codigo(self, anchor, texto, titulo=None):
        """Inserta un fragmento de código en Consolas 8,5 pt, fondo gris, sin interlineado.

        Cada línea es un párrafo; al final agrega un párrafo vacío de separación.
        """
        if titulo:
            anchor = self.parrafo(anchor, titulo)
        for linea in texto.rstrip('\n').split('\n'):
            p = self._nuevo_parrafo('Normal')
            pf = p.paragraph_format
            pf.space_after = Emu(0)
            pf.space_before = Emu(0)
            pf.line_spacing = 1.0
            self._run(p, linea if linea.strip() else ' ', fuente='Consolas', tam=Pt(8.5))
            sombra = OxmlElement('w:shd')
            sombra.set(qn('w:val'), 'clear')
            sombra.set(qn('w:color'), 'auto')
            sombra.set(qn('w:fill'), 'F2F2F2')
            p._p.get_or_add_pPr().append(sombra)
            anchor = self._insertar(anchor, p)
        return self.parrafo(anchor, '')


def reemplazar_texto(parrafo, nuevo: str):
    """Reemplaza el texto de un párrafo existente conservando el formato de
    su primer run (fuente, tamaño, negrita)."""
    runs = parrafo.runs
    if not runs:
        parrafo.add_run(nuevo)
        return
    runs[0].text = nuevo
    for r in runs[1:]:
        r.text = ''


def quitar_numeracion(parrafo):
    """Quita la numeración automática (numPr) de un párrafo, si la tiene."""
    pPr = parrafo._p.pPr
    if pPr is not None and pPr.numPr is not None:
        pPr.remove(pPr.numPr)
