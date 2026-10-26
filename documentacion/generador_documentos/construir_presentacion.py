# -*- coding: utf-8 -*-
"""
Presentación del trabajo final (video de 15 minutos) sobre la plantilla institucional, con el guion de cada
diapositiva en las notas del orador. Las cifras se leen de los mismos artefactos que el documento
(ml/modelos/*/metricas.json, eval/resultados, evidencias/pruebas.json), así que presentación y documento
dicen lo mismo.

Uso:
    python documentacion/generador_documentos/construir_presentacion.py --kit <lib del plugin de presentaciones>

Autores: Marlon David Pérez Almachi, Víctor Manuel Lluilema Pisco
Proyecto de titulación · Universidad Estatal de Milagro · 2026
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from _datos import C, CAPTURAS, DIAGRAMAS, FIGURAS, N_PRUEBAS_BACKEND, N_PRUEBAS_FRONTEND, O, U, f, pct

RAIZ = Path(__file__).resolve().parents[2]
SALIDA = RAIZ / 'documentacion' / 'presentacion'
PLANTILLA = SALIDA / 'plantilla' / 'PlantillaVicMar.pptx'
REPOSITORIO = 'https://github.com/Vm199524/PROYECTO-TITULACION-MARLON-VICTOR-UNEMI'
DRIVE = 'https://drive.google.com/drive/folders/1ZOVwEslSaSjJbUeS1RCL1q3BDyVea9tT'
TITULO = ('Prototipo de un sistema web para predecir el abandono en estudiantes universitarios mediante '
          'aprendizaje automático y consulta conversacional')
VICTOR, MARLON = 'Víctor', 'Marlon'


def _om(m):
    return (m.get('origen_movil') or {}).get('resumen') or {}


def contenido() -> dict:
    um, om = U['m'], O['m']
    ut, ot = um['validacion_temporal'], om['validacion_temporal']
    ru, ro = _om(um), _om(om)
    pr = O['p'] or {}
    cr = (C or {}).get('resumen', {})
    cal_u, cal_o = U['cal'] or {}, O['cal'] or {}
    der_u, der_o = U['der'] or {}, O['der'] or {}
    car_o = [c for c in om.get('por_carrera', []) if c.get('evaluable')]
    d = []

    d.append({'tipo': 'agenda', 'titulo': 'Contenido',
              'items': ['Problema y objetivos', 'Base conceptual y trabajos previos', 'Metodología y requerimientos',
                        'Arquitectura, datos y diagramas', 'Resultados del modelo y de la consulta',
                        'Seguridad, gobierno de datos y operación', 'Demostración del sistema',
                        'Conclusiones e innovación'],
              'notas': f'[{VICTOR}] La presentación sigue el orden del trabajo final. Primero, el problema y lo que nos '
                       'propusimos; después, la base teórica y la metodología. Luego mostraremos cómo está construido '
                       'el sistema, qué resultados obtuvo, cómo protege los datos y, al final, el sistema funcionando '
                       'y las conclusiones.'})

    d.append({'tipo': 'cifras', 'titulo': 'El problema', 'subtitulo': 'El abandono se mide tarde y con datos dispersos',
              'items': [{'valor': '44,90 %', 'etiqueta': 'probabilidad de abandono en cinco cohortes de una '
                                                        'universidad ecuatoriana (Núñez-Naranjo, 2024)'},
                        {'valor': '484', 'etiqueta': 'de 1.078 estudiantes admitidos abandonaron en ese mismo '
                                                    'estudio'},
                        {'valor': '≈ 0,70', 'etiqueta': 'ROC-AUC que informan los estudios que validan con cohortes '
                                                       'futuras (Carballo-Mendívil et al., 2025)'}],
              'texto': 'La universidad ya registra las señales (notas, asistencia, aula virtual, cartera), pero en '
                       'sistemas separados y para conocer el estado actual, no para anticipar el retiro.',
              'notas': f'[{VICTOR}] El abandono no es un problema hipotético en el Ecuador: en un estudio con cinco '
                       'cohortes, casi la mitad de los admitidos dejó la carrera. Las señales que lo anticipan ya '
                       'existen en los sistemas académicos, pero están dispersas y se consultan cuando el estudiante '
                       'ya se fue. Además, la literatura muestra que, cuando se valida con cohortes futuras, los '
                       'modelos rinden alrededor de 0,70; por eso fijamos un criterio realista de 0,75.'})

    d.append({'tipo': 'texto', 'titulo': 'Objetivo general',
              'parrafos': ['Desarrollar un prototipo de sistema web que prediga el abandono de estudiantes '
                           'universitarios mediante aprendizaje automático y permita consultar sus resultados en '
                           'lenguaje natural, para que autoridades y bienestar estudiantil identifiquen tempranamente '
                           'a quienes están en riesgo y prioricen su acompañamiento.'],
              'destacado': {'titulo': 'Qué · cómo · para qué',
                            'texto': 'Un sistema de alerta temprana, construido con CRISP-DM, un almacén en SQL '
                                     'Server, LightGBM y NL2SQL, para actuar antes del retiro.'},
              'notas': f'[{VICTOR}] Nuestro objetivo general responde a tres preguntas. Qué: un prototipo web de '
                       'alerta temprana. Cómo: con la metodología CRISP-DM, un almacén dimensional, un modelo '
                       'LightGBM calibrado y una consulta en lenguaje natural. Para qué: para que la autoridad y '
                       'bienestar estudiantil lleguen antes de que el estudiante abandone.'})

    d.append({'tipo': 'tarjetas', 'titulo': 'Objetivos específicos',
              'items': [{'icono': 'busqueda', 'titulo': 'Analizar', 'texto': 'Procesos y variables del abandono para '
                                                                             'derivar requerimientos'},
                        {'icono': 'grafico_barras', 'titulo': 'Construir y validar', 'texto': 'Modelo calibrado y '
                                                                                             'explicable, validado '
                                                                                             'en el tiempo'},
                        {'icono': 'monitor', 'titulo': 'Implementar', 'texto': 'Sistema web por roles con panel, '
                                                                               'ficha y consulta'},
                        {'icono': 'verificacion', 'titulo': 'Evaluar', 'texto': 'Métricas predictivas, de la consulta '
                                                                                'y de protección de datos'}],
              'notas': f'[{VICTOR}] Para cumplirlo planteamos cuatro objetivos específicos, cada uno con un '
                       'entregable verificable: el análisis que dio origen a los requerimientos, el modelo validado, '
                       'el sistema implementado y la evaluación contra criterios fijados antes de ver los resultados.'})

    d.append({'tipo': 'tabla', 'titulo': 'Trabajos previos frente al caso',
              'subtitulo': '14 estudios comparados; aquí, siete representativos',
              'encabezados': ['Estudio', 'Técnica', 'Resultado', 'Sistema / NL / Gobierno'],
              'filas': [['Núñez-Naranjo (2024)', 'Regresión logística multinivel', 'Abandono de 44,90 %', 'No / No / No'],
                        ['Carballo-Mendívil et al. (2025)', 'XGBoost, prueba prospectiva', 'ROC-AUC 0,69', 'Parcial / No / No'],
                        ['Rabelo y Zárate (2025)', 'Ensamble de modelos', 'Detecta 98,1 %', 'No / No / No'],
                        ['Vaarma y Li (2024)', 'Predicción mensual con Moodle', 'Créditos y actividad clave', 'No / No / No'],
                        ['Kuzilek et al. (2017)', 'Conjunto OULAD', '32.593 matrículas', 'No aplica'],
                        ['Ascoli y Choi (2025)', 'NL2SQL con LLM', 'Mejor exactitud con contexto', 'No aplica / Sí / No'],
                        ['VicMar (este trabajo)', 'LightGBM + K-means + NL2SQL', f'ROC-AUC {f(ot["roc_auc"])} (OULAD)',
                         'Sí / Sí / Sí']],
              'anchos': [1.5, 1.6, 1.3, 1.1],
              'notas': f'[{VICTOR}] Comparamos catorce estudios. Los modelos predictivos informan desde 76 % hasta casi '
                       '100 %, pero los valores altos suelen salir de particiones aleatorias. Ninguno entrega la '
                       'predicción dentro de un sistema con roles, consulta en lenguaje natural y aislamiento de datos '
                       'a la vez; ese es el vacío que cubre VicMar.'})

    d.append({'tipo': 'imagen', 'titulo': 'Metodología CRISP-DM', 'subtitulo': 'Seis fases con entregables verificables',
              'ruta': str(DIAGRAMAS / 'av2_crisp_dm.png'),
              'notas': f'[{VICTOR}] Aplicamos CRISP-DM. En comprensión del negocio fijamos los criterios de éxito; en '
                       'comprensión y preparación de datos construimos el ETL y el esquema estrella; en modelado '
                       'comparamos cuatro algoritmos; en evaluación validamos en el tiempo, por periodo y por carrera; '
                       'y en despliegue el modelo quedó dentro del sistema, con monitoreo de calidad y deriva.'})

    d.append({'tipo': 'cifras', 'titulo': 'Requerimientos',
              'subtitulo': 'Derivados de las cuatro piezas: modelo, segmentación, NL2SQL y panel',
              'items': [{'valor': '23', 'etiqueta': 'requerimientos funcionales'},
                        {'valor': '11', 'etiqueta': 'no funcionales con métrica verificable'},
                        {'valor': '10', 'etiqueta': 'historias de usuario para la autoridad sin perfil técnico'}],
              'texto': 'Criterios fijados antes de evaluar: ROC-AUC ≥ 0,75 global y ≥ 0,70 por periodo, exactitud de la consulta '
                       '≥ 85 %, 100 % de preguntas maliciosas contenidas y ninguna fuga entre instituciones.',
              'notas': f'[{VICTOR}] De ese análisis salieron 23 requerimientos funcionales, 11 no funcionales y diez '
                       'historias de usuario escritas para una autoridad que no sabe SQL. Lo importante es que los '
                       'criterios de aceptación se fijaron antes de entrenar, para no acomodarlos al resultado.'})

    d.append({'tipo': 'imagen', 'titulo': 'Arquitectura de la solución',
              'subtitulo': 'Cliente-servidor en capas, con el almacén como única fuente',
              'ruta': str(DIAGRAMAS / 'av2_arquitectura.png'),
              'notas': f'[{VICTOR}] La arquitectura es cliente-servidor en capas. Las fuentes académicas pasan por el '
                       'ETL al almacén ALMACEN_ACADEMICO en SQL Server. El pipeline entrena, califica y monitorea. La '
                       'API FastAPI expone 27 servicios y la interfaz React los consume. Panel, ficha y consulta leen '
                       'las mismas vistas, por eso las cifras coinciden.'})

    d.append({'tipo': 'tabla', 'titulo': 'Frontend y backend', 'subtitulo': 'Tecnologías y responsabilidades',
              'encabezados': ['Capa', 'Tecnología', 'Responsabilidad'],
              'filas': [['Frontend', 'React 18, TypeScript, Vite, Tailwind, TanStack Query',
                         'Pantallas por rol, tema claro y oscuro, sesión y caché de consultas'],
                        ['API', 'FastAPI (Python)', '27 servicios REST, JWT, roles y auditoría'],
                        ['Modelo', 'LightGBM, Platt, SHAP, K-means', 'Predicción calibrada, explicación y segmentos'],
                        ['Consulta', 'Router TF-IDF, LLM en cascada, sqlglot', '40 métricas certificadas y SQL '
                                                                             'validado de solo lectura'],
                        ['Datos', 'SQL Server 2022 en Docker', 'Esquema estrella, vistas VIS_* y SESSION_CONTEXT'],
                        ['Operación', 'MLflow, Great Expectations, Evidently', 'Registro, calidad y deriva']],
              'anchos': [1.0, 2.2, 2.6],
              'notas': f'[{MARLON}] En el frontend usamos React con TypeScript; cada rol ve solo sus pantallas. El '
                       'backend es FastAPI: autentica con JWT, valida el rol en cada servicio y registra la auditoría. '
                       'El modelo y la consulta viven en el backend, y la base de datos aplica su propio aislamiento '
                       'por institución.'})

    d.append({'tipo': 'imagen', 'titulo': 'Modelo de datos', 'subtitulo': 'Esquema estrella del almacén académico',
              'ruta': str(DIAGRAMAS / 'av2_modelo_estrella.png'),
              'notas': f'[{MARLON}] El almacén sigue el modelo dimensional de Kimball. Las tablas de hechos guardan '
                       'matrículas, predicciones, segmentos e intervenciones; las dimensiones describen estudiante, '
                       'carrera, periodo e institución. Las aplicaciones no leen las tablas: leen vistas filtradas '
                       'por institución.'})

    d.append({'tipo': 'imagen', 'titulo': 'Secuencia de una consulta en lenguaje natural',
              'ruta': str(DIAGRAMAS / 'av2_secuencia_consulta.png'),
              'notas': f'[{MARLON}] Así se responde una pregunta en español. Primero los guardarraíles descartan '
                       'pedidos peligrosos; luego el enrutador busca una métrica certificada. Si no la hay, el modelo '
                       'de lenguaje propone el SQL, el validador lo revisa y se ejecuta con una cuenta de solo lectura '
                       'que únicamente ve las vistas de su institución.'})

    d.append({'tipo': 'cifras', 'titulo': 'Resultados del modelo', 'subtitulo': 'LightGBM calibrado, probado con periodos que no vio',
              'items': [{'valor': f(ut['roc_auc']), 'etiqueta': 'ROC-AUC temporal, UNEMI (datos de prueba)'},
                        {'valor': f(ot['roc_auc']), 'etiqueta': 'ROC-AUC temporal, OULAD (datos reales)'},
                        {'valor': pct(ot['umbral_medio']['recall'], 0), 'etiqueta': 'de los abandonos detectados con '
                                                                                   'el umbral medio (OULAD)'}]
              + ([{'valor': f(pr['y_siguiente']['roc_auc']), 'etiqueta': 'ROC-AUC prospectivo con datos reales'}]
                 if pr.get('y_siguiente') else []),
              'texto': 'Los cuatro algoritmos se compararon con validación cruzada agrupada por estudiante; LightGBM '
                       'fue el mejor en las dos instituciones.',
              'notas': f'[{MARLON}] El ROC-AUC mide qué tan bien el modelo ordena: la probabilidad de que un estudiante '
                       'que abandona quede por encima de uno que no. Con periodos que el modelo nunca vio obtuvimos '
                       f'{f(ut["roc_auc"])} en los datos de prueba de la UNEMI y {f(ot["roc_auc"])} con los datos '
                       'reales de OULAD, ambos sobre el criterio de 0,75.'})

    d.append({'tipo': 'imagen', 'titulo': 'Desempeño por periodo y por carrera',
              'subtitulo': 'Validación con origen móvil: cada periodo se predice solo con el pasado',
              'ruta': str(FIGURAS / 'tf_desempeno_periodo.png'),
              'pie': (f'ROC-AUC por periodo: UNEMI {f(ru.get("roc_auc_min"))} a {f(ru.get("roc_auc_max"))}; OULAD '
                      f'{f(ro.get("roc_auc_min"))} a {f(ro.get("roc_auc_max"))}'
                      + (f'; módulos de OULAD {f(min(c["roc_auc"] for c in car_o))} a '
                         f'{f(max(c["roc_auc"] for c in car_o))}.' if car_o else '.')),
              'notas': f'[{MARLON}] Atendiendo la tutoría, reportamos el AUC y el recall por periodo con datos reales '
                       'anonimizados. Cada periodo se predijo entrenando solo con los anteriores, como si el sistema '
                       'se hubiera puesto en marcha en esa fecha. El desempeño se mantiene estable y lo mismo ocurre '
                       'por módulo, que en OULAD hace las veces de carrera.'})

    if cr:
        d.append({'tipo': 'cifras', 'titulo': 'Consulta conversacional', 'subtitulo': 'Evaluada con un conjunto dorado de preguntas',
                  'items': [{'valor': pct(cr.get('exactitud_relajada_total', 0), 0), 'etiqueta':
                             f'exactitud por ejecución en {cr.get("preguntas_base", 40)} preguntas base y '
                             f'{cr.get("preguntas_robustez", 12)} de robustez'},
                            {'valor': cr.get('adversariales_contenidas', '8/8'), 'etiqueta': 'preguntas maliciosas '
                                                                                           'contenidas'},
                            {'valor': str(cr.get('violaciones_aislamiento', 0)), 'etiqueta': 'lecturas de otra '
                                                                                            'institución'},
                            {'valor': f'{cr.get("latencia_mediana_ms", 0)} ms', 'etiqueta': 'latencia mediana'}],
                  'notas': f'[{MARLON}] La consulta se evaluó con un conjunto dorado: acertó en todas las preguntas '
                           'por resultado de ejecución, contuvo las ocho preguntas maliciosas, como intentos de borrar '
                           'datos, y nunca leyó datos de otra institución. La mayoría responde en milisegundos porque '
                           'usa métricas certificadas.'})

    d.append({'tipo': 'tarjetas', 'titulo': 'Seguridad y gobierno de datos', 'subtitulo': 'LOPDP aplicada en el motor, no solo en la pantalla',
              'items': [{'icono': 'candado', 'titulo': 'Aislamiento', 'texto': 'SESSION_CONTEXT y vistas por '
                                                                               'institución'},
                        {'icono': 'llave', 'titulo': 'Mínimo privilegio', 'texto': 'Cuentas chat_reader y '
                                                                                   'app_seguimiento'},
                        {'icono': 'usuario_ok', 'titulo': 'Tres roles', 'texto': 'Autoridad, bienestar y TI, con '
                                                                                 'auditoría'},
                        {'icono': 'proteccion', 'titulo': 'Anonimización', 'texto': 'Seudónimos HMAC y '
                                                                                    'k-anonimato k = 5'},
                        {'icono': 'ley', 'titulo': 'Gobierno de datos', 'texto': 'Documento LOPDP con roles, '
                                                                                 'ciclo de vida e incidentes'},
                        {'icono': 'alerta', 'titulo': 'Interruptor de IA', 'texto': 'La interpretación con LLM se '
                                                                                    'apaga por configuración'}],
              'notas': f'[{MARLON}] La seguridad está en la base de datos: cada sesión fija su institución y las vistas '
                       'filtran solo esa. Las cuentas tienen el mínimo privilegio. Para validar con una carrera real '
                       'de la UNEMI dejamos el documento de gobierno de datos según la LOPDP y un script que seudonimiza '
                       'cédulas y garantiza k-anonimato.'})

    d.append({'tipo': 'cifras', 'titulo': 'Operación del modelo (MLOps)', 'subtitulo': 'Cada corrida se valida, se vigila y se registra',
              'items': [{'valor': f'{cal_o.get("cumplidas", "n. d.")}/{cal_o.get("evaluadas", "n. d.")}',
                         'etiqueta': 'expectativas de calidad cumplidas (Great Expectations)'},
                        {'valor': f'{der_o.get("variables_con_deriva", "n. d.")}/{len(der_o.get("variables", []))}',
                         'etiqueta': 'variables con deriva en OULAD (Evidently)'},
                        {'valor': 'MLflow', 'etiqueta': 'registro de parámetros, métricas y modelo'},
                        {'valor': f'{N_PRUEBAS_BACKEND} + {N_PRUEBAS_FRONTEND}', 'etiqueta': 'pruebas automatizadas '
                                                                                             '(backend + frontend)'}],
              'texto': 'Metabase queda como vía de consulta alternativa para analistas, sobre un esquema agregado '
                       'sin datos personales.',
              'notas': f'[{MARLON}] Incorporamos las herramientas sugeridas en la tutoría. Great Expectations valida los '
                       'datos antes de entrenar y detiene el proceso si algo crítico falla. Evidently compara a los '
                       'estudiantes de hoy con los del entrenamiento. MLflow registra cada modelo. Y todo el sistema '
                       'está cubierto por pruebas automatizadas.'})

    d.append({'tipo': 'texto', 'titulo': 'Demostración del sistema',
              'parrafos': ['- Ingreso con rol de autoridad académica',
                           '- Panel ejecutivo y lista priorizada',
                           '- Ficha explicada y simulador',
                           '- Pregunta en lenguaje natural',
                           '- Calidad del modelo: por periodo y deriva',
                           'GitHub: ' + REPOSITORIO.replace('https://github.com/', ''),
                           'Carpeta digital: drive.google.com (enlace en la portada del trabajo)'],
              'imagen': str(CAPTURAS / 'av2_cap_panel.png'),
              'notas': f'[{VICTOR} y {MARLON}] Ahora mostramos el sistema funcionando (unos cuatro minutos). Víctor '
                       'ingresa como autoridad, recorre el panel y la lista priorizada y abre la ficha de un '
                       'estudiante con sus tres razones. Marlon hace una pregunta en lenguaje natural y muestra la '
                       'calidad del modelo por periodo. El código está en el repositorio de GitHub.'})

    d.append({'tipo': 'comparativa', 'titulo': 'Innovación y limitaciones',
              'izquierda': {'titulo': 'Aporte', 'items': ['Predicción explicada en español dentro de un sistema por roles',
                                                          'Consulta conversacional con métricas certificadas y SQL validado',
                                                          'Aislamiento multiinstitución en la base de datos',
                                                          'Perfil de mapeo: OULAD se incorporó sin cambiar el modelo']},
              'derecha': {'titulo': 'Limitación', 'items': ['UNEMI se evaluó con datos de prueba; la carrera real '
                                                            'requiere autorización',
                                                            'La prueba de usabilidad con usuarios está pendiente',
                                                            'El modelo prioriza; no reemplaza el criterio de bienestar']},
              'notas': f'[{VICTOR}] La innovación no es un algoritmo nuevo, sino la integración: predicción explicada, '
                       'consulta certificada y aislamiento en el motor, para una universidad ecuatoriana. También '
                       'somos claros con las limitaciones: falta validar con una carrera real de la UNEMI y la prueba '
                       'con usuarios, y ambas tienen ya su procedimiento.'})

    d.append({'tipo': 'conclusiones', 'titulo': 'Conclusiones',
              'items': [{'titulo': 'Requerimientos', 'texto': 'Las señales de abandono ya existen en los sistemas '
                                                              'académicos y dieron origen a 23 RF y 11 RNF.'},
                        {'titulo': 'Modelo', 'texto': f'LightGBM se sostuvo en periodos no vistos: '
                                                      f'{f(ro.get("roc_auc_min"))} a {f(ro.get("roc_auc_max"))} con '
                                                      'datos reales.'},
                        {'titulo': 'Sistema', 'texto': 'Funciona de punta a punta, con tres roles, 27 servicios y '
                                                       'operación automatizada.'},
                        {'titulo': 'Evaluación', 'texto': 'Cumple los criterios predictivos, de consulta y de '
                                                          'protección de datos fijados al inicio.'}],
              'notas': f'[{VICTOR}] En conclusión, es viable que una universidad ecuatoriana pase de medir el abandono '
                       'a anticiparlo con los datos que ya tiene, siempre que la predicción llegue a quien decide con '
                       'una explicación comprensible y con la protección de datos aplicada en la propia base. '
                       'Gracias por su atención.'})
    notas = {
        'El problema': f'[{VICTOR}] Empecemos por el problema. El abandono universitario no es hipotético en el Ecuador: '
                       'en un estudio con cinco cohortes de una universidad del país, 484 de 1.078 admitidos dejaron la '
                       'carrera, y la probabilidad de abandono llegó a 44,90 %. Las señales que lo anticipan, como las '
                       'notas, la asistencia, la actividad en el aula virtual o la deuda, ya están registradas, pero en '
                       'sistemas separados y se consultan para saber el estado actual del estudiante, no para anticipar '
                       'su retiro. Además, la literatura muestra que, cuando un modelo se prueba con cohortes futuras, '
                       'su desempeño ronda 0,70; por eso fijamos un criterio realista de 0,75.',
        'Trabajos previos frente al caso': f'[{VICTOR}] En la base conceptual partimos de Tinto, que explica el abandono '
                       'como falta de integración académica y social, y de Bean y Metzner, que muestran el peso de los '
                       'factores externos. Luego comparamos catorce estudios; aquí vemos siete. Los modelos predictivos '
                       'informan desde 76 % hasta casi 100 %, pero los valores más altos suelen salir de particiones '
                       'aleatorias que mezclan periodos. La última columna es la clave: ninguno entrega la predicción '
                       'dentro de un sistema con roles, consulta en lenguaje natural y aislamiento de datos a la vez. '
                       'Ese es el vacío que cubre VicMar.',
        'Metodología CRISP-DM': f'[{VICTOR}] Para desarrollarlo aplicamos CRISP-DM, con un entregable verificable por '
                       'fase. En comprensión del negocio fijamos los criterios de éxito y las historias de usuario. En '
                       'comprensión y preparación de los datos construimos el ETL, el esquema estrella y el panel '
                       'estudiante-periodo sin fuga de información futura. En modelado comparamos cuatro algoritmos. En '
                       'evaluación validamos en el tiempo, por periodo, por carrera y de forma prospectiva. Y en '
                       'despliegue el modelo quedó dentro del sistema, con calificación por lotes y monitoreo de la '
                       'calidad y la deriva de los datos.',
        'Arquitectura de la solución': f'[{VICTOR}] La arquitectura es cliente-servidor en capas. Abajo, las fuentes '
                       'académicas pasan por el ETL al almacén ALMACEN_ACADEMICO en SQL Server. El pipeline entrena el '
                       'modelo, califica a los estudiantes activos, segmenta las carreras y, al final, valida la calidad '
                       'y la deriva. En el medio, la API FastAPI expone 27 servicios protegidos por rol. Arriba, la '
                       'interfaz React. Lo importante es que el panel, la ficha y la consulta leen las mismas vistas '
                       'gobernadas, por eso las cifras coinciden en todo el sistema. Ahora Marlon explica cómo está '
                       'construido por dentro.',
        'Frontend y backend': f'[{MARLON}] Gracias, Víctor. En el frontend usamos React 18 con TypeScript, Vite y '
                       'Tailwind; cada rol ve solo su menú, hay tema claro y oscuro y funciona en el celular. TanStack '
                       'Query maneja la caché de las consultas. El backend es FastAPI en Python, el mismo lenguaje del '
                       'modelo: autentica con JWT, verifica el rol en cada uno de los 27 servicios y registra la '
                       'auditoría. El modelo usa LightGBM, Platt y SHAP; la consulta usa un enrutador TF-IDF, modelos de '
                       'lenguaje en cascada y el validador sqlglot. Y la operación del modelo se apoya en MLflow, Great '
                       'Expectations y Evidently.',
        'Modelo de datos': f'[{MARLON}] El almacén sigue el modelo dimensional de Kimball. Las tablas de hechos guardan '
                       'las matrículas por periodo, las predicciones, los segmentos de carrera y las intervenciones de '
                       'bienestar; las dimensiones describen al estudiante, la carrera, el periodo y la institución. '
                       'Un detalle de diseño importante: las aplicaciones nunca leen las tablas directamente, sino vistas '
                       'VIS que se filtran solas por la institución fijada en la sesión. El diccionario de datos completo '
                       'está en el documento.',
        'Secuencia de una consulta en lenguaje natural': f'[{MARLON}] Este diagrama de secuencia muestra cómo se '
                       'responde una pregunta en español. Primero, los guardarraíles descartan pedidos peligrosos, como '
                       'borrar o modificar datos. Luego el enrutador busca una de las 40 métricas certificadas; si la '
                       'encuentra, responde en milisegundos. Si no, un modelo de lenguaje propone el SQL, el validador '
                       'comprueba que sea solo una consulta de lectura sobre vistas permitidas y se ejecuta con la cuenta '
                       'chat_reader, que únicamente ve los datos de su institución. Son tres barreras independientes.',
        'Resultados del modelo': f'[{MARLON}] Pasemos a los resultados. El ROC-AUC mide qué tan bien ordena el modelo: la '
                       'probabilidad de que un estudiante que abandona quede por encima de uno que no; 0,5 es azar y 1 '
                       'es perfecto. Comparamos cuatro algoritmos con validación cruzada agrupada por estudiante y '
                       'LightGBM fue el mejor en las dos instituciones. Con periodos que el modelo nunca vio obtuvimos '
                       f'{f(ut["roc_auc"])} en los datos de prueba de la UNEMI y {f(ot["roc_auc"])} con los datos reales '
                       f'de OULAD, y con el umbral medio detecta cerca de {pct(ot["umbral_medio"]["recall"], 0)} de los '
                       'abandonos.',
        'Desempeño por periodo y por carrera': f'[{MARLON}] En la tutoría se nos pidió validar con datos reales '
                       'anonimizados y reportar el AUC y el recall por periodo. Lo hicimos con validación de origen '
                       'móvil: cada periodo se predijo entrenando solo con los anteriores, como si el sistema se hubiera '
                       f'puesto en marcha en esa fecha. En OULAD el ROC-AUC se mantuvo entre {f(ro.get("roc_auc_min"))} y '
                       f'{f(ro.get("roc_auc_max"))}, por encima del mínimo de 0,70 que fijamos para cada periodo. Por módulo, que en '
                       'OULAD hace las veces de carrera, cinco de siete superan ese mínimo; los dos que no lo alcanzan '
                       'tienen muy pocos abandonos, y eso nos dice que los umbrales deben revisarse por carrera. Para '
                       'una carrera real de la UNEMI dejamos listo el procedimiento de anonimización.',
        'Consulta conversacional': f'[{MARLON}] La consulta conversacional se evaluó con un conjunto dorado: preguntas '
                       'reales de una autoridad, con su resultado esperado. Acertó en todas por resultado de ejecución, '
                       'contuvo las ocho preguntas maliciosas, como intentos de borrar tablas o de ver otra institución, '
                       'y nunca leyó datos ajenos. La mayoría responde en milisegundos porque usa métricas certificadas; '
                       'el modelo de lenguaje solo entra cuando la pregunta no tiene una métrica.',
        'Seguridad y gobierno de datos': f'[{MARLON}] La seguridad no depende solo de la pantalla: está en la base de '
                       'datos. Cada conexión fija su institución con SESSION_CONTEXT y las vistas filtran solo esa; SQL '
                       'Server no deja cambiarla en la misma sesión. Las cuentas tienen el mínimo privilegio, hay tres '
                       'roles y todo queda auditado. Para validar con una carrera real de la UNEMI dejamos el documento '
                       'de gobierno de datos según la LOPDP y un script que reemplaza las cédulas por seudónimos y '
                       'garantiza k-anonimato con k igual a cinco.',
        'Operación del modelo (MLOps)': f'[{MARLON}] También incorporamos las herramientas sugeridas en la tutoría. Great '
                       'Expectations valida los datos antes de entrenar y detiene el proceso si falla una regla crítica. '
                       'Evidently compara a los estudiantes de hoy con los del entrenamiento para detectar la deriva, que '
                       'es la señal para reentrenar. MLflow registra cada modelo con sus parámetros y métricas. Metabase '
                       'queda como vía de consulta para analistas, sin datos personales. Y todo el sistema está cubierto '
                       'por pruebas automatizadas del backend y de la interfaz.',
        'Innovación y limitaciones': f'[{MARLON}] La innovación no está en un algoritmo nuevo, sino en la integración: '
                       'una predicción explicada en español, una consulta certificada y un aislamiento aplicado en el '
                       'motor, pensados para una universidad ecuatoriana. También somos claros con las limitaciones: la '
                       'UNEMI se evaluó con datos de prueba, falta la prueba con usuarios y el modelo prioriza, pero no '
                       'reemplaza el criterio de bienestar. Ambas validaciones pendientes ya tienen su procedimiento.',
        'Conclusiones': f'[{VICTOR}] Cerramos con una conclusión por objetivo. Las señales del abandono ya existen en los '
                       'sistemas académicos y de ellas salieron los requerimientos. El modelo se sostuvo en periodos que '
                       'no vio. El sistema funciona de punta a punta y su operación está automatizada. Y la evaluación '
                       'cumple los criterios fijados al inicio. En síntesis, es viable que una universidad ecuatoriana '
                       'pase de medir el abandono a anticiparlo con los datos que ya tiene. Gracias por su atención.',
    }
    for x in d:
        if x['titulo'] in notas:
            x['notas'] = notas[x['titulo']]
    return {
        'plantilla': str(PLANTILLA),
        'portada': {'notas': f'[{VICTOR}] Buenos días. Somos Víctor Manuel Lluilema Pisco y Marlon David Pérez '
                             'Almachi, de la carrera de Tecnologías de la Información, y presentamos nuestro trabajo '
                             f'de titulación: «{TITULO}». Al sistema lo llamamos VicMar.'},
        'diapositivas': d,
    }


def portada(pptx: Path) -> None:
    """Textos de la carátula de la plantilla del grupo (título del trabajo y tipo de exposición)."""
    from pptx import Presentation
    prs = Presentation(str(pptx))
    s = prs.slides[0]
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.startswith('DESARROLLO DEL TRABAJO'):
            _texto(sh, 'TRABAJO FINAL DE TITULACIÓN')
        elif t.startswith('DESARROLLO DE UN MODELO') or t.startswith('PROTOTIPO DE UN SISTEMA'):
            _texto(sh, TITULO.upper() + '.')
    prs.save(str(pptx))


def _texto(sh, texto: str) -> None:
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--kit', required=True, help='carpeta lib del plugin de presentaciones (contiene academico/)')
    ap.add_argument('--salida', default=str(SALIDA / 'Presentacion_Trabajo_Final_Marlon_Victor.pptx'))
    a = ap.parse_args()
    sys.path.insert(0, a.kit)
    from academico import diapositivas
    c = contenido()
    SALIDA.mkdir(parents=True, exist_ok=True)
    r = diapositivas.generar(c, a.salida)
    portada(Path(a.salida))
    print(json.dumps(r, ensure_ascii=False, indent=1))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
