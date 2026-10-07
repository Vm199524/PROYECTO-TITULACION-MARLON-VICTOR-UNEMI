# -*- coding: utf-8 -*-
"""
Reportes ejecutivos en Excel y PDF con la identidad del sistema.

Un reporte impreso debe reconocerse como parte del sistema: banda azul
institucional, línea de acento, indicadores en tarjetas, una «lectura del
periodo» con conclusiones generadas por reglas sobre los datos (nunca
inventadas), un gráfico y tablas con encabezado repetido. Formatos del
Ecuador: 1.234,5 y fechas dd/mm/aaaa. Las cédulas siempre van enmascaradas.

Autores: Marlon David Pérez Almachi, Víctor Manuel Lluilema Pisco
Proyecto de titulación · Universidad Estatal de Milagro · 2026
"""
from __future__ import annotations

import datetime as dt
import io
from typing import Any

AZUL = '002E45'
AZUL_MEDIO = '0B5C7E'
NARANJA = 'FF6900'
GRIS_900, GRIS_600, GRIS_200, GRIS_50 = '0F172A', '475569', 'E2E8F0', 'F8FAFC'
ROJO, AMBAR, VERDE = 'DC2626', 'D97706', '059669'
SISTEMA = 'VicMar · Alerta temprana de abandono'


def _num(v: Any, dec: int = 0) -> str:
    if v is None:
        return '—'
    texto = f'{float(v):,.{dec}f}'
    return texto.replace(',', '§').replace('.', ',').replace('§', '.')


def _pct(v: Any, dec: int = 1) -> str:
    return '—' if v is None else f'{_num(v, dec)} %'


def _seguro(v: Any) -> Any:
    """Evita la inyección de fórmulas en Excel (texto que empieza con =, +, - o @)."""
    if isinstance(v, str) and v[:1] in ('=', '+', '-', '@'):
        return "'" + v
    return v


def lectura(resumen: dict) -> list[str]:
    """Conclusiones del periodo generadas por reglas sobre las cifras (sin texto libre de IA)."""
    ind, conf = resumen['indicadores'], resumen['confiabilidad']
    frases = []
    if ind['evaluados']:
        frases.append(f"De {_num(ind['evaluados'])} estudiantes activos evaluados, {_num(ind['alto'])} "
                      f"({_pct(100 * ind['alto'] / ind['evaluados'])}) están en riesgo alto y {_num(ind['medio'])} en riesgo medio "
                      'de abandonar en el próximo periodo.')
    if resumen['riesgo_por_carrera']:
        c = resumen['riesgo_por_carrera'][0]
        frases.append(f"La carrera con más estudiantes en riesgo es {c['carrera']}: {_num(c['alto'])} en riesgo alto y "
                      f"{_num(c['medio'])} en riesgo medio de {_num(c['activos'])} activos.")
    if ind.get('carreras_prioritarias'):
        n = ind['carreras_prioritarias']
        frases.append(f"{_num(n)} {'carrera quedó' if n == 1 else 'carreras quedaron'} en el segmento de atención "
                      'prioritaria por su abandono histórico, rendimiento, asistencia y riesgo actual.')
    if conf.get('RECALL_MEDIO') is not None:
        frases.append(f"En la prueba temporal, las alertas de riesgo alto y medio detectaron el {_pct(100 * conf['RECALL_MEDIO'], 0)} de los "
                      f"abandonos reales, con una precisión del {_pct(100 * (conf.get('PRECISION_MEDIO') or 0), 0)} "
                      f"(ROC-AUC {_num(conf.get('ROC_AUC'), 3)}).")
    if ind.get('cobertura_alto') is not None:
        frases.append(f"El {_pct(ind['cobertura_alto'])} de los estudiantes en riesgo alto ya tiene una intervención registrada "
                      f"({_num(ind['intervenciones_abiertas'])} abiertas).")
    inst = resumen.get('institucion') or {}
    if inst.get('tipo_datos') == 'SINTETICO':
        frases.append('Atención: esta institución usa datos sintéticos de prueba; las cifras validan el funcionamiento del '
                      'sistema, no describen a estudiantes reales.')
    return frases


# ── Excel ────────────────────────────────────────────────────────────────
def _hoja_tabla(ws, titulo: str, columnas: list[str], filas: list[list], formatos: dict[int, str] | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    ws['A1'] = titulo
    ws['A1'].font = Font(bold=True, size=13, color=AZUL)
    borde = Border(bottom=Side(style='thin', color=GRIS_200))
    for j, c in enumerate(columnas, 1):
        celda = ws.cell(row=3, column=j, value=c)
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=AZUL)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, fila in enumerate(filas, 4):
        for j, v in enumerate(fila, 1):
            celda = ws.cell(row=i, column=j, value=_seguro(v))
            celda.border = borde
            if (formatos or {}).get(j):
                celda.number_format = formatos[j]
            if i % 2 == 0:
                celda.fill = PatternFill('solid', fgColor=GRIS_50)
    for j, c in enumerate(columnas, 1):
        largo = max([len(str(c))] + [len(str(f[j - 1])) for f in filas[:200]]) if filas else len(c)
        ws.column_dimensions[get_column_letter(j)].width = min(max(12, largo + 2), 60)
    ws.freeze_panes = 'A4'
    if filas:
        ws.auto_filter.ref = f'A3:{get_column_letter(len(columnas))}{3 + len(filas)}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = f'{SISTEMA} · Página &P de &N'


def excel_tabla(titulo: str, columnas: list[str], filas: list[list], sesion) -> bytes:
    """Resultado de una consulta en un Excel con el formato del sistema."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resultado'
    _hoja_tabla(ws, titulo[:120], columnas, filas)
    ws.cell(row=2, column=1, value=f'{sesion.institucion} · generado por {sesion.usuario} el '
                                   f'{dt.datetime.now().strftime("%d/%m/%Y %H:%M")}')
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def excel_ejecutivo(resumen: dict, alertas: list[dict], carreras: list[dict], sesion) -> bytes:
    """Reporte ejecutivo en Excel: resumen con gráfico nativo, alertas y carreras."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    ws.sheet_view.showGridLines = False
    for col in 'ABCDEFGH':
        ws.column_dimensions[col].width = 18
    ws.merge_cells('A1:H2')
    ws['A1'] = f'{SISTEMA.upper()} · REPORTE DE ALERTA TEMPRANA'
    ws['A1'].font = Font(bold=True, size=15, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=AZUL)
    ws['A1'].alignment = Alignment(vertical='center', indent=1)
    inst = resumen.get('institucion') or {}
    ws['A3'] = f"{inst.get('nombre', sesion.institucion)} · {inst.get('tipo_datos', '')} · generado por {sesion.usuario} " \
               f"el {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A3'].font = Font(color=GRIS_600, italic=True)
    ind, conf = resumen['indicadores'], resumen['confiabilidad']
    tarjetas = [('Evaluados', ind['evaluados'], '#,##0'), ('Riesgo alto', ind['alto'], '#,##0'),
                ('Riesgo medio', ind['medio'], '#,##0'), ('Abandono histórico (%)', ind['tasa_abandono_hist'], '0.0'),
                ('Carreras prioritarias', ind['carreras_prioritarias'], '0'),
                ('ROC-AUC temporal', conf.get('ROC_AUC'), '0.000'),
                ('Detección (recall)', conf.get('RECALL_MEDIO'), '0.0%'), ('Cobertura seguimiento (%)', ind.get('cobertura_alto'), '0.0')]
    for j, (etiqueta, valor, fmt) in enumerate(tarjetas, 1):
        ws.cell(row=5, column=j, value=etiqueta).font = Font(size=9, color=GRIS_600, bold=True)
        c = ws.cell(row=6, column=j, value=valor)
        c.font = Font(size=16, bold=True, color=AZUL)
        c.number_format = fmt
        ws.cell(row=7, column=j).fill = PatternFill('solid', fgColor=NARANJA)
    ws['A9'] = 'Lectura del periodo'
    ws['A9'].font = Font(bold=True, size=12, color=AZUL)
    for i, frase in enumerate(lectura(resumen), 10):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.cell(row=i, column=1, value='• ' + frase).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[i].height = 32
    fila0 = 18
    ws.cell(row=fila0, column=1, value='Carrera').font = Font(bold=True)
    ws.cell(row=fila0, column=2, value='Riesgo alto').font = Font(bold=True)
    ws.cell(row=fila0, column=3, value='Riesgo medio').font = Font(bold=True)
    for i, c in enumerate(resumen['riesgo_por_carrera'][:10], fila0 + 1):
        ws.cell(row=i, column=1, value=_seguro(c['carrera']))
        ws.cell(row=i, column=2, value=c['alto'])
        ws.cell(row=i, column=3, value=c['medio'])
    n = min(10, len(resumen['riesgo_por_carrera']))
    if n:
        grafico = BarChart()
        grafico.type, grafico.grouping, grafico.overlap = 'bar', 'stacked', 100
        grafico.title, grafico.height, grafico.width = 'Estudiantes en riesgo por carrera', 9, 18
        grafico.add_data(Reference(ws, min_col=2, max_col=3, min_row=fila0, max_row=fila0 + n), titles_from_data=True)
        grafico.set_categories(Reference(ws, min_col=1, min_row=fila0 + 1, max_row=fila0 + n))
        grafico.series[0].graphicalProperties.solidFill = ROJO
        grafico.series[1].graphicalProperties.solidFill = AMBAR
        ws.add_chart(grafico, f'E{fila0}')

    _hoja_tabla(wb.create_sheet('Alertas'), 'Estudiantes en riesgo (priorizados por probabilidad)',
                ['Estudiante', 'Cédula', 'Carrera', 'Probabilidad', 'Riesgo', 'Razón principal', 'Segunda razón'],
                [[a['nombre'], a['cedula'], a['carrera'], a['probabilidad'], a['riesgo'], a['factor_1'], a['factor_2']]
                 for a in alertas], {4: '0.0%'})
    _hoja_tabla(wb.create_sheet('Carreras'), 'Segmentación de carreras por perfil de riesgo',
                ['Carrera', 'Facultad', 'Segmento', 'Estudiantes', 'Abandono histórico', 'Promedio', 'Asistencia (%)',
                 'Riesgo medio', 'Activos en riesgo'],
                [[c['carrera'], c['facultad_nombre'], c['segmento'], c['estudiantes'], c['tasa_abandono_hist'], c['promedio_hist'],
                  c['asistencia_hist'], c['riesgo_promedio'], c['activos_en_riesgo']] for c in carreras],
                {5: '0.0%', 6: '0.00', 7: '0.0', 8: '0.0%'})
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────
def pdf_ejecutivo(resumen: dict, alertas: list[dict], carreras: list[dict], sesion) -> bytes:
    """Reporte ejecutivo en PDF (A4): portada con banda, indicadores, lectura, gráfico y tablas."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    c_azul, c_naranja = colors.HexColor('#' + AZUL), colors.HexColor('#' + NARANJA)
    c_gris, c_gris_claro = colors.HexColor('#' + GRIS_600), colors.HexColor('#' + GRIS_200)
    base = ParagraphStyle('base', fontName='Helvetica', fontSize=9.5, leading=13, textColor=colors.HexColor('#' + GRIS_900))
    titulo = ParagraphStyle('titulo', parent=base, fontName='Helvetica-Bold', fontSize=12.5, textColor=c_azul,
                            spaceBefore=8, spaceAfter=4)
    pequeno = ParagraphStyle('pequeno', parent=base, fontSize=8, leading=10, textColor=c_gris)
    celda = ParagraphStyle('celda', parent=base, fontSize=8, leading=10)
    inst = resumen.get('institucion') or {}
    generado = f"Generado por {sesion.usuario} el {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"

    def portada(canvas, doc):
        canvas.saveState()
        ancho, alto = A4
        if doc.page == 1:
            canvas.setFillColor(c_azul)
            canvas.rect(0, alto - 46 * mm, ancho, 46 * mm, stroke=0, fill=1)
            canvas.setFillColor(c_naranja)
            canvas.rect(0, alto - 47.5 * mm, ancho, 1.5 * mm, stroke=0, fill=1)
            canvas.setFillColor(colors.white)
            canvas.setFont('Helvetica-Bold', 10)
            canvas.drawString(18 * mm, alto - 14 * mm, SISTEMA.upper())
            canvas.setFont('Helvetica-Bold', 19)
            canvas.drawString(18 * mm, alto - 26 * mm, 'Reporte de alerta temprana de abandono')
            canvas.setFont('Helvetica', 10)
            canvas.drawString(18 * mm, alto - 34 * mm, f"{inst.get('nombre', sesion.institucion)} · "
                                                         f"{'datos sintéticos de prueba' if inst.get('tipo_datos') == 'SINTETICO' else 'datos reales anonimizados'}")
            canvas.drawString(18 * mm, alto - 40 * mm, generado)
        else:
            canvas.setFillColor(c_azul)
            canvas.rect(0, alto - 10 * mm, ancho, 10 * mm, stroke=0, fill=1)
            canvas.setFillColor(colors.white)
            canvas.setFont('Helvetica-Bold', 8.5)
            canvas.drawString(18 * mm, alto - 6.5 * mm, f'{SISTEMA.upper()} · REPORTE DE ALERTA TEMPRANA')
        canvas.setFillColor(c_gris)
        canvas.setFont('Helvetica', 7.5)
        canvas.drawString(18 * mm, 10 * mm, f'{SISTEMA} · {generado} · uso institucional, contiene datos personales protegidos')
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f'Página {doc.page}')
        canvas.restoreState()

    ind, conf = resumen['indicadores'], resumen['confiabilidad']
    tarjetas = [('ESTUDIANTES EVALUADOS', _num(ind['evaluados'])), ('RIESGO ALTO', _num(ind['alto'])),
                ('RIESGO MEDIO', _num(ind['medio'])), ('ABANDONO HISTÓRICO', _pct(ind['tasa_abandono_hist'])),
                ('ROC-AUC TEMPORAL', _num(conf.get('ROC_AUC'), 3)),
                ('ABANDONOS DETECTADOS', _pct(100 * conf['RECALL_MEDIO'], 0) if conf.get('RECALL_MEDIO') is not None else '—')]
    celdas = [[Paragraph(f'<font size=7 color="#{GRIS_600}"><b>{e}</b></font><br/><font size=15 color="#{AZUL}"><b>{v}</b></font>', base)
               for e, v in tarjetas[i:i + 3]] for i in (0, 3)]
    t_kpi = Table(celdas, colWidths=[58 * mm] * 3, rowHeights=[17 * mm] * 2)
    t_kpi.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), 0.5, c_gris_claro), ('INNERGRID', (0, 0), (-1, -1), 0.5, c_gris_claro),
                               ('LINEABOVE', (0, 0), (-1, 0), 2, c_naranja), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                               ('LEFTPADDING', (0, 0), (-1, -1), 8)]))
    historia = [Spacer(1, 38 * mm), Paragraph('Indicadores clave', titulo), t_kpi, Spacer(1, 4 * mm),
                Paragraph('Lectura del periodo', titulo)]
    lectura_t = Table([[Paragraph('• ' + f, base)] for f in lectura(resumen)], colWidths=[174 * mm])
    lectura_t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#FFF4EC')),
                                   ('LINEBEFORE', (0, 0), (0, -1), 2.5, c_naranja), ('LEFTPADDING', (0, 0), (-1, -1), 8)]))
    historia.append(lectura_t)

    datos = resumen['riesgo_por_carrera'][:10][::-1]
    if datos:
        d = Drawing(174 * mm, 78 * mm)
        g = HorizontalBarChart()
        g.x, g.y, g.width, g.height = 52 * mm, 6 * mm, 112 * mm, 66 * mm
        g.data = [[c['alto'] for c in datos], [c['medio'] for c in datos]]
        g.categoryAxis.categoryNames = [c['carrera'][:32] for c in datos]
        g.categoryAxis.labels.fontSize, g.valueAxis.labels.fontSize = 7, 7
        g.categoryAxis.style, g.barSpacing, g.groupSpacing = 'stacked', 0, 4
        g.bars[0].fillColor, g.bars[1].fillColor = colors.HexColor('#' + ROJO), colors.HexColor('#' + AMBAR)
        g.bars.strokeColor = None
        g.valueAxis.valueMin = 0
        d.add(g)
        d.add(String(52 * mm, 74 * mm, 'Rojo: riesgo alto · Ámbar: riesgo medio', fontSize=7, fillColor=c_gris))
        historia += [KeepTogether([Paragraph('Estudiantes en riesgo por carrera', titulo), d])]

    def tabla(encabezado, filas, anchos):
        t = Table([[Paragraph(f'<b>{h}</b>', ParagraphStyle('h', parent=celda, textColor=colors.white)) for h in encabezado]] +
                  [[Paragraph(str(v), celda) for v in f] for f in filas], colWidths=anchos, repeatRows=1)
        estilo = [('BACKGROUND', (0, 0), (-1, 0), c_azul), ('LINEBELOW', (0, 1), (-1, -1), 0.3, c_gris_claro),
                  ('VALIGN', (0, 0), (-1, -1), 'TOP')]
        estilo += [('BACKGROUND', (0, i), (-1, i), colors.HexColor('#' + GRIS_50)) for i in range(2, len(filas) + 1, 2)]
        t.setStyle(TableStyle(estilo))
        return t

    historia += [Paragraph('Carreras por segmento de riesgo', titulo),
                 tabla(['Carrera', 'Segmento', 'Abandono hist.', 'Riesgo medio', 'En riesgo'],
                       [[c['carrera'], c['segmento'], _pct(100 * c['tasa_abandono_hist']), _pct(100 * c['riesgo_promedio']),
                         _num(c['activos_en_riesgo'])] for c in carreras],
                       [62 * mm, 38 * mm, 25 * mm, 25 * mm, 24 * mm])]
    historia += [Paragraph('Estudiantes con mayor riesgo', titulo),
                 tabla(['Estudiante', 'Carrera', 'Prob.', 'Razón principal'],
                       [[f"{a['nombre']}<br/><font size=7 color='#{GRIS_600}'>{a['cedula']}</font>", a['carrera'],
                         _pct(100 * a['probabilidad']), a['factor_1'] or '—'] for a in alertas[:25]],
                       [48 * mm, 40 * mm, 16 * mm, 70 * mm]),
                 Spacer(1, 3 * mm),
                 Paragraph('Las probabilidades son estimaciones del modelo predictivo para priorizar el acompañamiento; no '
                           'deben usarse para sancionar ni negar beneficios. El listado completo está en el reporte en Excel.', pequeno)]
    salida = io.BytesIO()
    SimpleDocTemplate(salida, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                      title='Reporte de alerta temprana', author='Marlon David Pérez Almachi; Víctor Manuel Lluilema Pisco'
                      ).build(historia, onFirstPage=portada, onLaterPages=portada)
    return salida.getvalue()
